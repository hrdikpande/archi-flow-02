"""S6: F-52 exports (Excel/PDF, background over the row limit) and F-54 onboarding."""

import io
from datetime import date

import openpyxl
import pytest
from django.utils import timezone

from apps.accounts.permissions import PermissionDeniedError
from apps.accounts.roles import ProjectRole
from apps.automation.models import Flag, FlagKind
from apps.core.exceptions import ConflictError, DomainError
from apps.exports import services as exports
from apps.exports.models import ExportState
from apps.notifications.models import Notification
from apps.outbox.models import OutboxEvent
from apps.projects.models import Project
from apps.register import onboarding
from apps.register.models import OnboardingState, RevisionState, Sheet, SheetState
from apps.register.onboarding_template import build_template

from .workflow_fixtures import new_sheet, ready_upload_for


@pytest.fixture
def owner_user(make_user):
    from apps.accounts.roles import OrgRole

    return make_user(OrgRole.OWNER)


def read_xlsx(data: bytes) -> list[list]:
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


class TestExports:
    def test_register_excel_with_filters(self, team, fake_storage):
        a = new_sheet(team, "A-101", title='=HYPERLINK("evil")')
        new_sheet(team, "A-102")
        Flag.objects.create(
            organisation=team.org,
            entity_type="sheet",
            entity_id=str(a.pk),
            kind=FlagKind.LATE,
            rule="A-02",
            sheet=a,
            project=team.project,
            raised_at=timezone.now(),
            data={"days": 2},
        )
        job = exports.start(team.pl, "register", "xlsx", {"late": "1"}, team.project)
        assert job.state == ExportState.DONE and job.row_count == 1
        rows = read_xlsx(fake_storage.objects[job.file.storage_key])
        header = rows.index(
            [
                "Sheet",
                "Title",
                "Disc",
                "Stage",
                "Package",
                "Owner",
                "Checker",
                "Planned",
                "State",
                "Rev",
                "Flags",
            ]
        )
        assert rows[header + 1][0] == "A-101"
        assert rows[header + 1][1].startswith("'=")  # formula neutralised
        assert rows[header + 1][-1] == "Late 2 d"

    def test_permissions_follow_the_list(self, team, fake_storage):
        with pytest.raises(PermissionDeniedError):
            exports.start(team.outsider, "register", "xlsx", {}, team.project)
        with pytest.raises(DomainError):
            exports.start(team.pl, "register", "xlsx", {}, None)

    def test_large_export_runs_in_background_and_emails(
        self, team, fake_storage, settings, django_capture_on_commit_callbacks
    ):
        settings.EXPORT_SYNC_MAX_ROWS = 1
        new_sheet(team, "A-101")
        new_sheet(team, "A-102")
        with django_capture_on_commit_callbacks(execute=True):
            job = exports.start(team.pl, "register", "xlsx", {}, team.project)
        assert job.state == ExportState.QUEUED
        job.refresh_from_db()
        assert job.state == ExportState.DONE and job.row_count == 2
        note = Notification.objects.get(user=team.pl, event="export_ready")
        assert note.link == f"/exports/{job.pk}"

    def test_pdf_and_other_kinds(self, team, fake_storage):
        new_sheet(team, "A-101")
        job = exports.start(team.pl, "register", "pdf", {}, team.project)
        assert fake_storage.objects[job.file.storage_key][:5] == b"%PDF-"
        for kind in ("transmittal_log", "late_at_risk", "workload", "timesheets"):
            job = exports.start(team.pl, kind, "xlsx", {}, team.project)
            assert job.state == ExportState.DONE, kind

    def test_api_returns_download_url(self, team, client_for, fake_storage):
        new_sheet(team, "A-101")
        res = client_for(team.pl).post(
            "/api/v1/exports",
            {"kind": "register", "format": "xlsx", "project": str(team.project.pk), "filters": {}},
            format="json",
        )
        assert res.status_code == 201 and res.json()["download_url"]
        other = client_for(team.outsider).get(f"/api/v1/exports/{res.json()['id']}")
        assert other.status_code == 404


def workbook(team, *, code="2026-050", extra_revision=None, bad=False) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(build_template()))
    project = wb["Project"]
    values = {
        "Code": code,
        "Name": "Raipur OPD",
        "Client": "State Health Society",
        "Type": "hospital",
        "Start date": "01-Jan-2026",
        "CAD versions": "Revit 2024",
        "Project lead email": team.pl.email,
    }
    for row in project.iter_rows(min_row=2):
        if row[0].value in values:
            row[1].value = values[row[0].value]
    wb["Stages"].append(
        ["1", "Concept", "40", "01-Jan-2026", "31-Mar-2026", "", "", "", "", "deliverables_issued"]
    )
    wb["Stages"].append(
        ["2", "Working drawings", "60", "01-Apr-2026", "", "", "", "FA,GFC", "yes", "active"]
    )
    wb["Team"].append([team.designer.email, "Designer", ""])
    wb["Team"].append([team.dl.email, "DL", "A"])
    wb["Packages"].append(["Arch GFC set", "A", "30-Nov-2026"])
    wb["Sheets"].append(
        [
            "A-101",
            "Ground Floor Plan",
            "A",
            "Working drawings",
            "Arch GFC set",
            team.designer.email,
            team.dl.email,
            "15-Nov-2026",
            "",
            "",
        ]
    )
    wb["Sheets"].append(
        ["A-102", "First Floor Plan", "A", "2", "", team.designer.email, "", "20-Nov-2026", "", ""]
    )
    if bad:
        wb["Sheets"].append(
            ["Q-1", "Bad", "Z", "Nowhere", "", "stranger@x.example", "", "31-Feb-2026", "", ""]
        )
    wb["Revisions"].append(
        ["A-101", "P01", "FA", "10-Feb-2026", "A-101 P01.pdf", "superseded", "First"]
    )
    wb["Revisions"].append(
        ["A-101", "C01", "GFC", "12-Mar-2026", "A-101 C01.pdf", "issued", "For site"]
    )
    if extra_revision:
        wb["Revisions"].append(extra_revision)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestOnboarding:
    def test_load_live_project_without_notifications(self, team, admin):
        job = onboarding.start(admin, name="raipur.xlsx", content=workbook(team))
        assert job.report["errors"] == []
        assert job.report["missing_files"] == ["a-101 c01.pdf", "a-101 p01.pdf"]
        with pytest.raises(ConflictError):
            onboarding.commit(admin, job)
        uploads = [ready_upload_for(admin, n, sha=n) for n in ("A-101 P01.pdf", "A-101 C01.pdf")]
        job = onboarding.attach_files(admin, job, [u.pk for u in uploads])
        assert job.report["ready"] is True
        events_before = OutboxEvent.objects.count()
        notes_before = Notification.objects.count()
        job = onboarding.commit(admin, job)
        assert job.state == OnboardingState.COMMITTED
        project = Project.objects.get(code="2026-050")
        assert project.stage_plan_final is True
        assert project.memberships.filter(user=team.pl, role=ProjectRole.PL).exists()
        assert project.memberships.get(user=team.dl).discipline.code == "A"
        a101 = Sheet.objects.get(project=project, number="A-101")
        assert a101.state == SheetState.ISSUED and a101.current_revision.code == "C01"
        revs = {r.code: r for r in a101.revisions.all()}
        assert revs["P01"].state == RevisionState.SUPERSEDED and revs["P01"].is_historic
        assert revs["C01"].issued_at.date() == date(2026, 3, 12)
        assert revs["P01"].superseded_at.date() == date(2026, 3, 12)
        assert Sheet.objects.get(project=project, number="A-102").state == SheetState.PLANNED
        assert OutboxEvent.objects.count() == events_before  # nothing emitted
        assert Notification.objects.count() == notes_before
        with pytest.raises(ConflictError):
            onboarding.commit(admin, job)  # only once

    def test_validation_lists_every_problem(self, team, admin):
        job = onboarding.start(
            admin,
            name="bad.xlsx",
            content=workbook(
                team, bad=True, extra_revision=["A-101", "C02", "XYZ", "", "", "issued", ""]
            ),
        )
        messages = [e["message"] for e in job.report["errors"]]
        assert any("Unknown discipline “Z”" in m for m in messages)
        assert any("Stage “Nowhere”" in m for m in messages)
        assert any("31-Feb-2026" in m for m in messages)
        assert any("Unknown purpose “XYZ”" in m for m in messages)
        assert any("Issued date is required" in m for m in messages)
        assert any("exactly one revision marked issued" in m for m in messages)
        assert job.report["ready"] is False

    def test_existing_code_and_admin_only(self, team, admin):
        job = onboarding.start(admin, name="dup.xlsx", content=workbook(team, code="2026-014"))
        assert any("already exists" in e["message"] for e in job.report["errors"])
        with pytest.raises(PermissionDeniedError):
            onboarding.start(team.pl, name="x.xlsx", content=workbook(team))


class TestProjectArchive:
    def test_full_project_export_json_and_files(
        self, team, fake_storage, owner_user, django_capture_on_commit_callbacks
    ):
        import json
        import zipfile

        from .workflow_fixtures import draft_with_pdf

        sheet = new_sheet(team, "A-101")
        rev = draft_with_pdf(team, sheet)
        fake_storage.objects[rev.file.storage_key] = b"%PDF-1.4 test"
        with django_capture_on_commit_callbacks(execute=True):
            job = exports.start(owner_user, "project_archive", "zip", {}, team.project)
        job.refresh_from_db()
        assert job.state == ExportState.DONE and job.row_count == 1
        zf = zipfile.ZipFile(io.BytesIO(fake_storage.objects[job.file.storage_key]))
        data = json.loads(zf.read("project.json"))
        assert data["format"] == "drawingops-project-export/1"
        assert [s["number"] for s in data["sheets"]] == ["A-101"]
        assert data["files"][0]["sha256"] == rev.file.sha256
        assert zf.read(data["files"][0]["path"]) == b"%PDF-1.4 test"

    def test_archive_needs_owner_or_admin_and_zip(self, team, fake_storage):
        with pytest.raises(PermissionDeniedError):
            exports.start(team.pl, "project_archive", "zip", {}, team.project)
        with pytest.raises(DomainError):
            exports.start(team.pl, "register", "zip", {}, team.project)
