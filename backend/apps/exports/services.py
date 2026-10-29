"""F-52 exports: Excel and PDF of any list with its current filters. Up to
EXPORT_SYNC_MAX_ROWS rows are built during the request; larger exports run as a background job
and the person gets an email with a link (the link opens the app, which issues a fresh,
short-lived download URL, so no long-lived file URL is ever emailed)."""

import hashlib
import io
from typing import Any
from zoneinfo import ZoneInfo

from django.conf import settings
from django.db import transaction
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.accounts.models import User
from apps.accounts.permissions import Action, require
from apps.audit.services import record
from apps.core.exceptions import DomainError
from apps.core.pdf import render_pdf
from apps.files import storage
from apps.files.models import FileObject
from apps.notifications.services import notify
from apps.projects.models import Project

from . import archive
from .builders import BUILDERS, Table
from .models import ExportFormat, ExportJob, ExportKind, ExportState

MIME: dict[str, str] = {
    ExportFormat.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ExportFormat.PDF: "application/pdf",
    ExportFormat.ZIP: "application/zip",
}


def _safe(value: Any) -> Any:
    """Neutralise spreadsheet formulas in user-entered text (CSV/Excel injection)."""
    if isinstance(value, str) and len(value) > 1 and value[0] in "=+-@\t\r":
        return f"'{value}"
    return value


def to_xlsx(table: Table) -> tuple[bytes, int]:
    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = table.title[:31].replace("/", "-").replace(":", " ") or "Export"
    ws.append([table.title])
    ws["A1"].font = Font(bold=True, size=13)
    if table.subtitle:
        ws.append([table.subtitle])
    ws.append([])
    ws.append(table.columns)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEF1F5")
    count = 0
    for row in table.rows:
        ws.append([_safe(v) for v in row])
        count += 1
    for i, width in enumerate(table.widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), count


def to_pdf(table: Table, firm: str, tz: str, filters: str) -> tuple[bytes, int]:

    rows = list(table.rows)
    html = render_to_string(
        "exports/table.html",
        {
            "title": table.title,
            "subtitle": table.subtitle,
            "columns": table.columns,
            "rows": rows,
            "count": len(rows),
            "firm": firm,
            "filters": filters,
            "generated": timezone.now().astimezone(ZoneInfo(tz)).strftime("%d-%b-%Y %H:%M"),
        },
    )
    return render_pdf(html), len(rows)


def _filters_text(filters: dict[str, Any]) -> str:
    parts = []
    for k, v in filters.items():
        if v in (None, "", [], False):
            continue
        parts.append(f"{k}={','.join(map(str, v)) if isinstance(v, list) else v}")
    return "; ".join(parts)


def start(
    user: User, kind: str, fmt: str, filters: dict[str, Any], project: Project | None
) -> ExportJob:
    archive_job = kind == ExportKind.PROJECT_ARCHIVE
    if not archive_job and kind not in BUILDERS:
        raise DomainError("invalid_kind", "Unknown export.")
    if archive_job != (fmt == ExportFormat.ZIP) or fmt not in ExportFormat.values:
        raise DomainError("invalid_format", "Lists export to Excel or PDF; a full project to zip.")
    if archive_job:
        if project is None:
            raise DomainError("project_required", "Choose a project to export.")
        require(user, Action.EXPORT_AUDIT)  # N-06 includes the project's audit trail
        count = archive.count(project)
    else:
        count = BUILDERS[kind](user, project, filters).count()  # checks permission
    job = ExportJob.objects.create(
        organisation_id=user.org_pk,
        user=user,
        kind=kind,
        format=fmt,
        filters=filters,
        project=project,
        row_count=count,
    )
    record(
        organisation_id=user.org_pk,
        actor=user,
        verb="export.requested",
        entity_type="exports.exportjob",
        entity_id=job.pk,
        entity_label=f"{kind}.{fmt} ({count} rows)",
        after={"filters": filters},
        project_id=project.pk if project else None,
    )
    if count <= settings.EXPORT_SYNC_MAX_ROWS and not archive_job:
        return build(job.pk)
    from .tasks import build_export

    transaction.on_commit(lambda: build_export.delay(str(job.pk)))
    return job


def build(job_id: Any) -> ExportJob:
    job = ExportJob.objects.select_related("user", "project", "organisation").get(pk=job_id)
    if job.state == ExportState.DONE:
        return job
    job.state = ExportState.RUNNING
    job.save(update_fields=["state", "updated_at"])
    try:
        if job.kind == ExportKind.PROJECT_ARCHIVE:
            return _build_archive(job)
        table = BUILDERS[job.kind](job.user, job.project, job.filters).table()
        org = job.organisation
        if job.format == ExportFormat.XLSX:
            data, count = to_xlsx(table)
        else:
            data, count = to_pdf(table, org.name, org.timezone, _filters_text(job.filters))
        sha = hashlib.sha256(data).hexdigest()
        stamp = timezone.now().astimezone(ZoneInfo(org.timezone)).strftime("%Y%m%d-%H%M")
        name = f"{table.title.replace('·', '-').replace('/', '-')} {stamp}.{job.format}"
        # The file name is the last key segment: storage that ignores response-header
        # overrides (SeaweedFS) still downloads it under the right name.
        key = f"exports/{job.organisation_id}/{job.pk}/{_key_name(name)}"
        storage.put_bytes(key, data, MIME[job.format])
        job.file, _ = FileObject.objects.get_or_create(
            organisation_id=job.organisation_id,
            sha256=sha,
            defaults={
                "storage_key": key,
                "size": len(data),
                "mime": MIME[job.format],
                "original_name": name,
            },
        )
        job.row_count = count
        job.state = ExportState.DONE
    except DomainError as exc:
        job.state, job.error = ExportState.FAILED, exc.detail
    except Exception as exc:
        job.state, job.error = ExportState.FAILED, repr(exc)[:2000]
        job.finished_at = timezone.now()
        job.save()
        raise
    job.finished_at = timezone.now()
    job.save()
    return job


def _build_archive(job: ExportJob) -> ExportJob:
    import tempfile

    if job.project is None:
        raise DomainError("project_required", "The project no longer exists.")
    stamp = timezone.now().astimezone(ZoneInfo(job.organisation.timezone)).strftime("%Y%m%d-%H%M")
    name = f"{job.project.code} full export {stamp}.zip"
    with tempfile.TemporaryFile() as tmp:
        files = archive.write_archive(job.project, tmp)
        size = tmp.tell()
        tmp.seek(0)
        sha = hashlib.sha256()
        while chunk := tmp.read(1 << 20):
            sha.update(chunk)
        tmp.seek(0)
        key = f"exports/{job.organisation_id}/{job.pk}/{_key_name(name)}"
        storage.put_file(key, tmp, MIME[ExportFormat.ZIP])
    job.file, _ = FileObject.objects.get_or_create(
        organisation_id=job.organisation_id,
        sha256=sha.hexdigest(),
        defaults={
            "storage_key": key,
            "size": size,
            "mime": MIME[ExportFormat.ZIP],
            "original_name": name,
        },
    )
    job.row_count = files
    job.state = ExportState.DONE
    job.finished_at = timezone.now()
    job.save()
    return job


def _key_name(name: str) -> str:
    """A storage-safe version of a download name: letters, digits, dot, dash, underscore."""
    import re

    return re.sub(r"[^\w.-]+", "_", name.strip()).strip("_")[:150] or "export"


def notify_ready(job: ExportJob) -> None:
    ok = job.state == ExportState.DONE
    notify(
        job.user,
        "export_ready",
        title=f"Your export is {'ready' if ok else 'not ready: it failed'} ({job.row_count} rows)",
        body="Open the link to download it." if ok else job.error,
        link=f"/exports/{job.pk}",
        dedupe_key=f"export:{job.pk}",
    )


def download_url(job: ExportJob, origin: str | None) -> str | None:
    if job.state != ExportState.DONE or job.file is None:
        return None
    return storage.presign_get(
        job.file.storage_key, job.file.original_name, origin=origin, attachment=True
    )
