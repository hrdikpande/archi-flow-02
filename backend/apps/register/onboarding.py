"""F-54 onboarding of a live project: project, stages, team, packages, sheets and past revisions
from the template workbook plus the PDFs it names.

Validation runs on upload and after every attached PDF; commit is all-or-nothing. Historic
revisions import as `issued`/`superseded` with their original dates and `is_historic=True`.
Nothing is emitted to the outbox, so no notifications or transmittal emails go out during
import (03 F-54: "nothing triggers notifications during import"); audit rows are written.
"""

import io
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from typing import Any
from zoneinfo import ZoneInfo

import openpyxl
from django.db import transaction
from django.utils import timezone

from apps.accounts.models import User
from apps.accounts.permissions import Action, require
from apps.accounts.roles import ProjectRole
from apps.audit.services import record
from apps.contacts.models import ExternalOrg
from apps.core.exceptions import ConflictError, DomainError
from apps.core.models import Discipline, OrgSettings
from apps.files.models import FileObject, Upload, UploadState
from apps.projects.models import (
    Package,
    Project,
    ProjectMembership,
    ProjectType,
    Stage,
    StageState,
)

from .imports import MAX_BYTES, _cell_text, parse_date
from .models import (
    OnboardingJob,
    OnboardingState,
    RevisionState,
    Sheet,
    SheetRevision,
    SheetState,
)
from .naming import sheet_number_problems
from .onboarding_template import SHEETS

ROLE_NAMES = {
    "pl": ProjectRole.PL,
    "project lead": ProjectRole.PL,
    "dl": ProjectRole.DL,
    "discipline lead": ProjectRole.DL,
    "designer": ProjectRole.DESIGNER,
    "senior reviewer": ProjectRole.SENIOR_REVIEWER,
    "senior_reviewer": ProjectRole.SENIOR_REVIEWER,
    "sr": ProjectRole.SENIOR_REVIEWER,
    "dc": ProjectRole.DC,
    "document controller": ProjectRole.DC,
    "viewer": ProjectRole.VIEWER,
}
STAGE_STATES = {s.value: s.value for s in StageState} | {
    s.label.lower(): s.value for s in StageState
}


def _key(header: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", header.lower()).strip("_")


def parse_workbook(name: str, content: bytes) -> dict[str, Any]:
    if len(content) > MAX_BYTES:
        raise DomainError("too_large", "The workbook is larger than 5 MB.")
    if not name.lower().endswith(".xlsx"):
        raise DomainError("unsupported", "Upload the onboarding template as an .xlsx file.")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise DomainError("unreadable", f"{name} could not be read as an Excel file.") from exc
    missing = [s for s in SHEETS if s not in wb.sheetnames]
    if missing:
        raise DomainError("missing_sheets", f"The workbook is missing: {', '.join(missing)}.")
    data: dict[str, Any] = {}
    for sheet, headers in SHEETS.items():
        rows = [[_cell_text(c) for c in r] for r in wb[sheet].iter_rows(values_only=True)]
        rows = [r for r in rows if any(c for c in r)]
        if sheet == "Project":
            data["project"] = {
                _key(r[0]): (r[1] if len(r) > 1 else "") for r in rows[1:] if r and r[0]
            }
            continue
        keys = [_key(h) for h in headers]
        data[_key(sheet)] = [
            {"row": i, **{k: (r[j] if j < len(r) else "") for j, k in enumerate(keys)}}
            for i, r in enumerate(rows[1:], start=2)
        ]
    return data


class _Report:
    def __init__(self) -> None:
        self.errors: list[dict[str, Any]] = []

    def add(self, sheet: str, row: int | None, message: str) -> None:
        self.errors.append({"sheet": sheet, "row": row, "message": message})


def _date(report: _Report, sheet: str, row: int, label: str, text: str) -> date | None:
    try:
        return parse_date(text or "")
    except ValueError:
        report.add(sheet, row, f"{label} “{text}” is not a date.")
        return None


def validate(job: OnboardingJob) -> dict[str, Any]:
    """Check the whole workbook against the firm's set-up; store and return the report."""
    org_id = job.organisation_id
    d = job.data
    report = _Report()
    p = d.get("project", {})
    code, name = (p.get("code") or "").strip(), (p.get("name") or "").strip()
    if not code:
        report.add("Project", None, "Code is required.")
    elif Project.objects.filter(organisation_id=org_id, code__iexact=code).exists():
        report.add("Project", None, f"A project with code {code} already exists.")
    if not name:
        report.add("Project", None, "Name is required.")
    for label in ("start_date", "end_date"):
        _date(report, "Project", 0, label.replace("_", " ").capitalize(), p.get(label, ""))
    users = {
        u.email.lower(): u for u in User.objects.filter(organisation_id=org_id, is_active=True)
    }
    lead = (p.get("project_lead_email") or "").strip().lower()
    if lead and lead not in users:
        report.add("Project", None, f"Project lead {lead} has not signed in to DrawingOps yet.")

    disciplines: dict[str, Discipline] = {}
    for d_ in Discipline.objects.filter(organisation_id=org_id):
        disciplines[d_.code.lower()] = d_
        disciplines[d_.name.lower()] = d_

    stage_keys: set[str] = set()
    fee_total = Decimal(0)
    for s in d.get("stages", []):
        if not s.get("name"):
            report.add("Stages", s["row"], "Name is required.")
            continue
        try:
            seq = int(s.get("seq") or 0)
        except ValueError:
            seq = 0
        if seq <= 0:
            report.add("Stages", s["row"], "Seq must be a whole number from 1.")
        try:
            fee_total += Decimal(s.get("fee") or "0")
        except InvalidOperation:
            report.add("Stages", s["row"], f"Fee % “{s.get('fee')}” is not a number.")
        for label in ("planned_start", "planned_end", "actual_start", "actual_end"):
            _date(
                report, "Stages", s["row"], label.replace("_", " ").capitalize(), s.get(label, "")
            )
        if s.get("state") and s["state"].lower() not in STAGE_STATES:
            report.add("Stages", s["row"], f"Unknown stage state “{s['state']}”.")
        stage_keys |= {s["name"].lower(), str(seq), f"{seq} {s['name']}".lower()}
    if not d.get("stages"):
        report.add("Stages", None, "Add the project's stages.")
    if fee_total and fee_total != 100:
        report.add("Stages", None, f"Fee % adds up to {fee_total}%, not 100%.")

    team_emails = {lead} if lead else set()
    for t in d.get("team", []):
        email = (t.get("email") or "").strip().lower()
        if email not in users:
            report.add("Team", t["row"], f"{email or 'Email'} has not signed in to DrawingOps yet.")
        if (t.get("role") or "").strip().lower() not in ROLE_NAMES:
            report.add("Team", t["row"], f"Unknown role “{t.get('role')}”.")
        if t.get("discipline") and t["discipline"].lower() not in disciplines:
            report.add("Team", t["row"], f"Unknown discipline “{t['discipline']}”.")
        team_emails.add(email)

    packages: set[str] = set()
    for pk in d.get("packages", []):
        if not pk.get("name"):
            report.add("Packages", pk["row"], "Name is required.")
            continue
        packages.add(pk["name"].lower())
        if pk.get("discipline") and pk["discipline"].lower() not in disciplines:
            report.add("Packages", pk["row"], f"Unknown discipline “{pk['discipline']}”.")
        _date(report, "Packages", pk["row"], "Planned date", pk.get("planned_date", ""))

    numbers: dict[str, int] = {}
    probe = Project(organisation_id=org_id, code=code or "X")
    for s in d.get("sheets", []):
        number = (s.get("sheet_number") or "").strip()
        if not number or not s.get("title"):
            report.add("Sheets", s["row"], "Sheet number and title are required.")
            continue
        if number in numbers:
            report.add("Sheets", s["row"], f"Sheet {number} appears more than once.")
        numbers[number] = s["row"]
        disc = disciplines.get((s.get("discipline") or "").lower())
        if disc is None:
            report.add("Sheets", s["row"], f"Unknown discipline “{s.get('discipline')}”.")
        else:
            for problem in sheet_number_problems(probe, disc, number):
                report.add("Sheets", s["row"], problem)
        if (s.get("stage") or "").lower() not in stage_keys:
            report.add("Sheets", s["row"], f"Stage “{s.get('stage')}” is not in the Stages sheet.")
        if s.get("package") and s["package"].lower() not in packages:
            report.add(
                "Sheets", s["row"], f"Package “{s['package']}” is not in the Packages sheet."
            )
        for role in ("owner", "checker"):
            email = (s.get(role) or "").strip().lower()
            if email and email not in team_emails:
                report.add(
                    "Sheets", s["row"], f"{role.capitalize()} {email} is not on the Team sheet."
                )
        _date(report, "Sheets", s["row"], "Planned issue date", s.get("planned_issue_date", ""))

    purposes = {
        p["code"]
        for p in OrgSettings.objects.get(organisation_id=org_id).purpose_codes
        if p.get("active", True)
    }
    issued_per_sheet: dict[str, int] = {}
    codes: set[tuple[str, str]] = set()
    needed_files: set[str] = set()
    for r in d.get("revisions", []):
        number = (r.get("sheet_number") or "").strip()
        if number not in numbers:
            report.add("Revisions", r["row"], f"Sheet {number or '?'} is not in the Sheets sheet.")
        code_ = (r.get("revision") or "").strip().upper()
        if not re.match(r"^[A-Z]{1,2}\d{1,3}$", code_):
            report.add(
                "Revisions", r["row"], f"Revision “{r.get('revision')}” is not like P01 or C02."
            )
        if (number, code_) in codes:
            report.add("Revisions", r["row"], f"{number} {code_} appears more than once.")
        codes.add((number, code_))
        if (r.get("purpose") or "").upper() not in purposes:
            report.add("Revisions", r["row"], f"Unknown purpose “{r.get('purpose')}”.")
        if not (r.get("issued_date") or "").strip():
            report.add("Revisions", r["row"], "Issued date is required.")
        else:
            _date(report, "Revisions", r["row"], "Issued date", r["issued_date"])
        status = (r.get("status") or "").strip().lower()
        if status not in ("issued", "superseded"):
            report.add("Revisions", r["row"], "Status must be issued or superseded.")
        elif status == "issued":
            issued_per_sheet[number] = issued_per_sheet.get(number, 0) + 1
        if r.get("file_name"):
            needed_files.add(r["file_name"].strip().lower())
    for number in {(r.get("sheet_number") or "").strip() for r in d.get("revisions", [])}:
        n = issued_per_sheet.get(number, 0)
        if n != 1:
            report.add(
                "Revisions",
                None,
                f"{number} needs exactly one revision marked issued (the current one); it has {n}.",
            )
    missing = sorted(needed_files - set(job.files))
    result = {
        "errors": report.errors,
        "missing_files": missing,
        "counts": {
            "stages": len(d.get("stages", [])),
            "team": len(d.get("team", [])),
            "packages": len(d.get("packages", [])),
            "sheets": len(d.get("sheets", [])),
            "revisions": len(d.get("revisions", [])),
            "files": len(job.files),
        },
        "ready": not report.errors and not missing,
    }
    job.report = result
    job.save(update_fields=["report", "updated_at"])
    return result


def start(actor: User, *, name: str, content: bytes) -> OnboardingJob:
    require(actor, Action.MANAGE_ORG_SETTINGS)
    job = OnboardingJob.objects.create(
        organisation_id=actor.org_pk,
        created_by=actor,
        original_name=name,
        data=parse_workbook(name, content),
    )
    validate(job)
    return job


def attach_files(actor: User, job: OnboardingJob, upload_ids: list[Any]) -> OnboardingJob:
    """PDFs uploaded through the normal pipeline (hash, page check) are matched by file name."""
    require(actor, Action.MANAGE_ORG_SETTINGS)
    _draft(job)
    uploads = Upload.objects.filter(
        pk__in=upload_ids, uploaded_by=actor, state=UploadState.READY
    ).select_related("file")
    files = dict(job.files)
    for u in uploads:
        if u.file_id:
            files[u.original_name.strip().lower()] = str(u.file_id)
    job.files = files
    job.save(update_fields=["files", "updated_at"])
    validate(job)
    return job


def _draft(job: OnboardingJob) -> None:
    if job.state != OnboardingState.DRAFT:
        raise ConflictError("committed", "This onboarding has already been loaded.")


def _at_noon(day: date, tz: ZoneInfo) -> datetime:
    return datetime.combine(day, time(12, 0), tzinfo=tz)


@transaction.atomic
def commit(actor: User, job: OnboardingJob) -> OnboardingJob:
    require(actor, Action.MANAGE_ORG_SETTINGS)
    job = OnboardingJob.objects.select_for_update().get(pk=job.pk)
    _draft(job)
    report = validate(job)
    if not report["ready"]:
        raise ConflictError(
            "not_ready",
            "Fix the listed problems and upload the missing PDFs first.",
            errors=report["errors"][:50],
            missing_files=report["missing_files"][:50],
        )
    org = actor.org
    tz = ZoneInfo(org.timezone)
    d = job.data
    p = d["project"]
    users = {u.email.lower(): u for u in User.objects.filter(organisation=org, is_active=True)}
    disciplines: dict[str, Discipline] = {}
    for d_ in Discipline.objects.filter(organisation=org):
        disciplines[d_.code.lower()] = d_
        disciplines[d_.name.lower()] = d_

    client = None
    if p.get("client"):
        client, _ = ExternalOrg.objects.get_or_create(
            organisation=org,
            name__iexact=p["client"],
            defaults={"name": p["client"], "kind": "client"},
        )
    ptype = (p.get("type") or "").lower()
    project = Project.objects.create(
        organisation=org,
        code=p["code"].strip(),
        name=p["name"].strip(),
        client=client,
        type=ptype if ptype in ProjectType.values else ProjectType.OTHER,
        location=p.get("location", ""),
        start_date=parse_date(p.get("start_date", "")),
        end_date=parse_date(p.get("end_date", "")),
        cad_versions=[v.strip() for v in (p.get("cad_versions") or "").split(",") if v.strip()],
    )

    stages: dict[str, Stage] = {}
    fee_total = Decimal(0)
    for s in d["stages"]:
        seq = int(s["seq"])
        fee = Decimal(s.get("fee") or "0")
        fee_total += fee
        stage = Stage.objects.create(
            organisation=org,
            project=project,
            seq=seq,
            name=s["name"],
            fee_pct=fee,
            planned_start=parse_date(s.get("planned_start", "")),
            planned_end=parse_date(s.get("planned_end", "")),
            actual_start=parse_date(s.get("actual_start", "")),
            actual_end=parse_date(s.get("actual_end", "")),
            state=STAGE_STATES.get((s.get("state") or "planned").lower(), StageState.PLANNED),
            requires_senior_qc=(s.get("senior_qc") or "").lower() in ("yes", "y", "true", "1"),
            required_purposes=[
                x.strip().upper()
                for x in (s.get("required_purposes") or "").split(",")
                if x.strip()
            ]
            or ["FA", "GFC"],
        )
        for k in (s["name"].lower(), str(seq), f"{seq} {s['name']}".lower()):
            stages[k] = stage
    project.stage_plan_final = fee_total == 100
    project.save(update_fields=["stage_plan_final", "updated_at"])

    lead = (p.get("project_lead_email") or "").strip().lower()
    team_rows = [{"email": lead, "role": "pl", "discipline": ""}] if lead else []
    for t in [*team_rows, *d.get("team", [])]:
        user = users[t["email"].strip().lower()]
        role = ROLE_NAMES[t["role"].strip().lower()]
        disc = disciplines.get((t.get("discipline") or "").lower())
        ProjectMembership.objects.get_or_create(
            organisation=org,
            project=project,
            user=user,
            role=role,
            defaults={"discipline": disc if role == ProjectRole.DL else None},
        )

    packages: dict[str, Package] = {}
    for pk in d.get("packages", []):
        packages[pk["name"].lower()] = Package.objects.create(
            organisation=org,
            project=project,
            name=pk["name"],
            discipline=disciplines.get((pk.get("discipline") or "").lower()),
            planned_date=parse_date(pk.get("planned_date", "")),
        )

    sheets: dict[str, Sheet] = {}
    for s in d["sheets"]:
        sheets[s["sheet_number"].strip()] = Sheet.objects.create(
            organisation=org,
            project=project,
            stage=stages[s["stage"].lower()],
            package=packages.get((s.get("package") or "").lower()),
            discipline=disciplines[s["discipline"].lower()],
            number=s["sheet_number"].strip(),
            title=s["title"],
            owner=users.get((s.get("owner") or "").lower()),
            checker=users.get((s.get("checker") or "").lower()),
            planned_issue_date=parse_date(s.get("planned_issue_date", "")),
            cad_link=s.get("cad_link", ""),
            notes=s.get("notes", ""),
        )

    by_sheet: dict[str, list[dict[str, Any]]] = {}
    for r in d.get("revisions", []):
        by_sheet.setdefault(r["sheet_number"].strip(), []).append(r)
    files = {k: FileObject.objects.get(pk=v, organisation=org) for k, v in job.files.items()}
    revisions = 0
    for number, rows in by_sheet.items():
        sheet = sheets[number]
        rows.sort(key=lambda r: parse_date(r["issued_date"]) or date.min)
        current = None
        for seq, r in enumerate(rows, start=1):
            issued_at = _at_noon(parse_date(r["issued_date"]) or date.today(), tz)
            later = rows[seq] if seq < len(rows) else None
            status = r["status"].strip().lower()
            file = files.get((r.get("file_name") or "").strip().lower())
            revision = SheetRevision.objects.create(
                organisation=org,
                sheet=sheet,
                code=r["revision"].strip().upper(),
                seq=seq,
                state=RevisionState.ISSUED if status == "issued" else RevisionState.SUPERSEDED,
                stage=sheet.stage,
                purpose=r["purpose"].strip().upper(),
                file=file,
                original_name=file.original_name if file else "",
                change_summary=(r.get("change_summary") or "")[:500],
                issued_at=issued_at,
                issued_by=actor,
                superseded_at=_at_noon(parse_date(later["issued_date"]) or date.today(), tz)
                if status == "superseded" and later
                else (issued_at if status == "superseded" else None),
                is_historic=True,
            )
            revisions += 1
            if status == "issued":
                current = revision
        sheet.current_revision = current  # validation guarantees exactly one issued revision
        sheet.state = SheetState.ISSUED
        sheet.state_changed_at = timezone.now()
        sheet.save(update_fields=["current_revision", "state", "state_changed_at", "updated_at"])

    job.state = OnboardingState.COMMITTED
    job.project = project
    job.committed_at = timezone.now()
    job.save(update_fields=["state", "project", "committed_at", "updated_at"])
    record(
        organisation_id=org.pk,
        actor=actor,
        verb="project.onboarded",
        entity_type="register.onboardingjob",
        entity_id=job.pk,
        entity_label=f"{project.code}: {len(sheets)} sheets, {revisions} revisions",
        after=job.report.get("counts"),
        project_id=project.pk,
    )
    return job
