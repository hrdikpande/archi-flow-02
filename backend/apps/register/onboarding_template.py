"""Template workbook for the register import (F-05) and project onboarding (F-54)."""

import io

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SHEETS: dict[str, list[str]] = {
    "Project": ["Field", "Value"],
    "Stages": [
        "Seq",
        "Name",
        "Fee %",
        "Planned start",
        "Planned end",
        "Actual start",
        "Actual end",
        "Required purposes",
        "Senior QC",
        "State",
    ],
    "Team": ["Email", "Role", "Discipline"],
    "Packages": ["Name", "Discipline", "Planned date"],
    "Sheets": [
        "Sheet number",
        "Title",
        "Discipline",
        "Stage",
        "Package",
        "Owner",
        "Checker",
        "Planned issue date",
        "CAD link",
        "Notes",
    ],
    "Revisions": [
        "Sheet number",
        "Revision",
        "Purpose",
        "Issued date",
        "File name",
        "Status",
        "Change summary",
    ],
}
PROJECT_FIELDS = [
    "Code",
    "Name",
    "Client",
    "Type",
    "Location",
    "Start date",
    "End date",
    "CAD versions",
    "Project lead email",
]
INSTRUCTIONS = [
    "DrawingOps import template",
    "",
    "Register import (F-05): fill in the 'Sheets' sheet only, "
    "then upload it on the project's Register.",
    "Onboarding a live project (F-54): fill in every sheet, "
    "then upload it with the PDFs in Admin › Onboarding.",
    "",
    "Dates: DD-Mon-YYYY (23-Oct-2026), YYYY-MM-DD or DD/MM/YYYY.",
    "Discipline: the code (A, M, E…) or name from Admin › Disciplines.",
    "Stage: the stage name or number from the project's stage plan.",
    "Owner / Checker / Team email: people must have signed in to DrawingOps once.",
    "Team role: PL, DL, Designer, Senior reviewer, DC or Viewer.",
    "Revisions: Status is 'issued' (current) or 'superseded'. "
    "File name must match a PDF you upload.",
    "Purpose: a purpose code from Admin › Organisation (e.g. FA, GFC). Never write 'IFC'.",
]


def build_template() -> bytes:
    wb = openpyxl.Workbook()
    info = wb.active
    info.title = "Instructions"
    for line in INSTRUCTIONS:
        info.append([line])
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 110
    header_fill = PatternFill("solid", fgColor="F1F3F6")
    for name, headers in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        for i, _ in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=i)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            ws.column_dimensions[get_column_letter(i)].width = 22
        ws.freeze_panes = "A2"
        if name == "Project":
            for field in PROJECT_FIELDS:
                ws.append([field, ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
