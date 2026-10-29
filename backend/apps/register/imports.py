"""Excel/CSV sheet-register import (F-05, I-01).

upload → column mapping (auto-guessed, editable) → preview with row errors (fix or skip) →
commit (all-or-nothing) → undo within 24 h if no imported sheet has moved since.
"""

import csv
import io
import re
from collections import Counter
from datetime import date, datetime, timedelta
from typing import Any

import openpyxl
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from apps.accounts.models import User
from apps.accounts.permissions import Action, require
from apps.audit.services import record, record_change, snapshot
from apps.core.exceptions import ConflictError, DomainError
from apps.core.models import Discipline
from apps.files.validators import validate_cad_link
from apps.outbox.services import emit
from apps.projects.models import Package, Project, ProjectMembership, Stage

from .models import ImportJob, ImportState, Sheet, SheetState
from .naming import sheet_number_problems

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10_000
UNDO_WINDOW = timedelta(hours=24)

FIELDS: dict[str, dict[str, Any]] = {
    "number": {
        "label": "Sheet number",
        "required": True,
        "synonyms": [
            "sheet number",
            "sheet no",
            "drawing number",
            "drawing no",
            "dwg no",
            "number",
            "no",
            "sheet",
        ],
    },
    "title": {
        "label": "Title",
        "required": True,
        "synonyms": ["title", "sheet title", "drawing title", "description", "name"],
    },
    "discipline": {
        "label": "Discipline",
        "required": True,
        "synonyms": ["discipline", "disc", "trade", "service"],
    },
    "stage": {"label": "Stage", "required": True, "synonyms": ["stage", "fee stage", "phase"]},
    "package": {"label": "Package", "required": False, "synonyms": ["package", "set", "pack"]},
    "owner": {
        "label": "Owner (designer)",
        "required": False,
        "synonyms": ["owner", "designer", "drafter", "assigned to", "drawn by"],
    },
    "checker": {
        "label": "Checker",
        "required": False,
        "synonyms": ["checker", "checked by", "reviewer"],
    },
    "planned_issue_date": {
        "label": "Planned issue date",
        "required": False,
        "synonyms": [
            "planned issue date",
            "planned date",
            "due date",
            "due",
            "target date",
            "issue date",
        ],
    },
    "cad_link": {
        "label": "CAD file location",
        "required": False,
        "synonyms": ["cad link", "cad file", "file path", "path", "location"],
    },
    "notes": {"label": "Notes", "required": False, "synonyms": ["notes", "remarks", "comments"]},
}

DATE_FORMATS = ("%d-%b-%Y", "%d-%b-%y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d %b %Y")


# --- Parsing --------------------------------------------------------------------------------


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_file(name: str, content: bytes) -> tuple[list[str], list[list[str]]]:
    if len(content) > MAX_BYTES:
        raise DomainError("too_large", "The file is larger than 5 MB. Split the register.")
    lower = name.lower()
    if lower.endswith(".xlsx"):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise DomainError("unreadable", f"{name} could not be read as an Excel file.") from exc
        ws = wb["Sheets"] if "Sheets" in wb.sheetnames else wb.active
        rows = [[_cell_text(c) for c in row] for row in ws.iter_rows(values_only=True)]
    elif lower.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    else:
        raise DomainError("unsupported", "Upload an .xlsx or .csv file.")
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        raise DomainError("empty", f"{name} has no rows.")
    headers, body = rows[0], rows[1:]
    if len(body) > MAX_ROWS:
        raise DomainError("too_many_rows", f"{name} has more than {MAX_ROWS:,} rows.")
    width = len(headers)
    body = [(r + [""] * width)[:width] for r in body]
    return headers, body


def _norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def guess_mapping(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    used: set[int] = set()
    normalized = [_norm(h) for h in headers]
    for field, spec in FIELDS.items():
        for synonym in spec["synonyms"]:
            for idx, header in enumerate(normalized):
                if idx not in used and header == synonym:
                    mapping[field] = idx
                    used.add(idx)
                    break
            if field in mapping:
                break
    return mapping


# --- Validation -----------------------------------------------------------------------------


def parse_date(text: str) -> date | None:
    text = text.strip()
    if not text:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


class _Lookups:
    def __init__(self, project: Project) -> None:
        self.disciplines: dict[str, Discipline] = {}
        for d in Discipline.objects.filter(organisation_id=project.organisation_id):
            self.disciplines[d.code.lower()] = d
            self.disciplines[d.name.lower()] = d
        self.stages: dict[str, Stage] = {}
        for s in Stage.objects.filter(project=project):
            self.stages[s.name.lower()] = s
            self.stages[str(s.seq)] = s
            self.stages[f"{s.seq} {s.name}".lower()] = s
        self.packages = {p.name.lower(): p for p in Package.objects.filter(project=project)}
        self.members: dict[str, User] = {}
        for m in ProjectMembership.objects.filter(project=project).select_related("user"):
            self.members[m.user.email.lower()] = m.user
            if m.user.full_name:
                self.members[m.user.full_name.lower()] = m.user
        self.existing = set(Sheet.objects.filter(project=project).values_list("number", flat=True))


def _row_values(job: ImportJob, index: int, row: list[str]) -> dict[str, str]:
    values = {f: (row[i] if i < len(row) else "") for f, i in job.mapping.items()}
    values.update(job.overrides.get(str(index), {}))
    return {k: (v or "").strip() for k, v in values.items()}


def validate(job: ImportJob) -> dict[str, Any]:
    """Row-level preview. Never writes. Returns {rows: [...], counts: {...}}."""
    project = job.project
    look = _Lookups(project)
    missing = [
        FIELDS[f]["label"] for f, s in FIELDS.items() if s["required"] and f not in job.mapping
    ]
    values_by_row = {i: _row_values(job, i, r) for i, r in enumerate(job.rows)}
    numbers = Counter(
        v.get("number", "") for i, v in values_by_row.items() if i not in set(job.skip_rows)
    )
    results = []
    for i, values in values_by_row.items():
        errors: list[str] = []
        if i in set(job.skip_rows):
            results.append(
                {"row": i + 2, "index": i, "status": "skipped", "values": values, "errors": []}
            )
            continue
        for field, spec in FIELDS.items():
            if spec["required"] and field in job.mapping and not values.get(field):
                errors.append(f"{spec['label']} is empty.")
        number = values.get("number", "")
        discipline = look.disciplines.get(values.get("discipline", "").lower())
        if values.get("discipline") and discipline is None:
            errors.append(f"Unknown discipline “{values['discipline']}”.")
        if number:
            if numbers[number] > 1:
                errors.append(f"Sheet {number} appears more than once in the file.")
            if number in look.existing:
                errors.append(f"Sheet {number} already exists in this project.")
            if discipline is not None:
                errors.extend(sheet_number_problems(project, discipline, number))
        if values.get("stage") and look.stages.get(values["stage"].lower()) is None:
            errors.append(f"Unknown stage “{values['stage']}”. Set up the stage plan first.")
        if values.get("package") and look.packages.get(values["package"].lower()) is None:
            errors.append(f"Unknown package “{values['package']}”. Create it first.")
        for key in ("owner", "checker"):
            if values.get(key) and look.members.get(values[key].lower()) is None:
                errors.append(f"{values[key]} is not on this project's team.")
        try:
            parse_date(values.get("planned_issue_date", ""))
        except ValueError:
            errors.append(
                f"Bad date “{values['planned_issue_date']}”. Use DD-Mon-YYYY, e.g. 23-Oct-2026."
            )
        if values.get("cad_link"):
            try:
                validate_cad_link(values["cad_link"])
            except ValidationError as exc:
                errors.append(exc.messages[0])
        results.append(
            {
                "row": i + 2,
                "index": i,
                "status": "error" if errors else "ok",
                "values": values,
                "errors": errors,
            }
        )
    counts = Counter(r["status"] for r in results)
    return {
        "missing_columns": missing,
        "rows": results,
        "counts": {
            "ok": counts["ok"],
            "error": counts["error"],
            "skipped": counts["skipped"],
            "total": len(results),
        },
    }


# --- Services -------------------------------------------------------------------------------


@transaction.atomic
def start_import(actor: User, project: Project, *, name: str, content: bytes) -> ImportJob:
    require(actor, Action.EDIT_SHEETS, project=project)
    headers, rows = parse_file(name, content)
    job = ImportJob.objects.create(
        organisation_id=project.organisation_id,
        created_by=actor,
        project=project,
        original_name=name[:255],
        headers=headers,
        rows=rows,
        mapping=guess_mapping(headers),
    )
    job.report = validate(job)
    job.save(update_fields=["report"])
    return job


def _editable(job: ImportJob) -> None:
    if job.state != ImportState.PREVIEW:
        raise ConflictError("not_editable", "This import has already been committed.")


@transaction.atomic
def update_import(
    actor: User,
    job: ImportJob,
    *,
    mapping: dict[str, int] | None = None,
    skip_rows: list[int] | None = None,
    overrides: dict[str, dict[str, str]] | None = None,
) -> ImportJob:
    require(actor, Action.EDIT_SHEETS, project=job.project)
    job = ImportJob.objects.select_for_update().get(pk=job.pk)
    _editable(job)
    if mapping is not None:
        unknown = set(mapping) - set(FIELDS)
        bad_cols = [v for v in mapping.values() if not 0 <= int(v) < len(job.headers)]
        if unknown or bad_cols:
            raise DomainError("invalid_mapping", "The column mapping refers to unknown columns.")
        job.mapping = {k: int(v) for k, v in mapping.items()}
    if skip_rows is not None:
        job.skip_rows = sorted({int(i) for i in skip_rows if 0 <= int(i) < len(job.rows)})
    if overrides is not None:
        job.overrides = {
            str(int(i)): {k: str(v) for k, v in vals.items() if k in FIELDS}
            for i, vals in overrides.items()
            if 0 <= int(i) < len(job.rows)
        }
    job.report = validate(job)
    job.save()
    return job


@transaction.atomic
def commit_import(actor: User, job: ImportJob) -> ImportJob:
    """All-or-nothing: any remaining error row blocks the whole import (F-05)."""
    require(actor, Action.EDIT_SHEETS, project=job.project)
    job = ImportJob.objects.select_for_update(of=("self",)).select_related("project").get(pk=job.pk)
    _editable(job)
    report = validate(job)
    if report["missing_columns"] or report["counts"]["error"]:
        job.report = report
        job.save(update_fields=["report"])
        raise ConflictError(
            "import_has_errors",
            f"{report['counts']['error']} rows still have errors. Fix or skip them, then import.",
            counts=report["counts"],
            missing_columns=report["missing_columns"],
        )
    look = _Lookups(job.project)
    created: list[Sheet] = []
    for item in report["rows"]:
        if item["status"] != "ok":
            continue
        v = item["values"]
        sheet = Sheet(
            organisation_id=job.organisation_id,
            created_by=actor,
            project=job.project,
            number=v["number"],
            title=v["title"],
            discipline=look.disciplines[v["discipline"].lower()],
            stage=look.stages[v["stage"].lower()],
            package=look.packages.get(v.get("package", "").lower()) if v.get("package") else None,
            owner=look.members.get(v.get("owner", "").lower()) if v.get("owner") else None,
            checker=look.members.get(v.get("checker", "").lower()) if v.get("checker") else None,
            planned_issue_date=parse_date(v.get("planned_issue_date", "")),
            cad_link=v.get("cad_link", ""),
            notes=v.get("notes", ""),
        )
        created.append(sheet)
    Sheet.objects.bulk_create(created)
    for sheet in created:
        record_change(sheet, actor=actor, verb="sheet.imported", before=None, after=snapshot(sheet))
    now = timezone.now()
    job.state = ImportState.COMMITTED
    job.report = report
    job.created_sheet_ids = [str(s.pk) for s in created]
    job.committed_at = now
    job.undo_until = now + UNDO_WINDOW
    job.save()
    record(
        organisation_id=job.organisation_id,
        actor=actor,
        verb="import.committed",
        entity_type="register.importjob",
        entity_id=job.pk,
        entity_label=f"{job.original_name} → {len(created)} sheets",
        project_id=job.project_id,
    )
    emit(
        "import.committed",
        {"import_id": str(job.pk), "project_id": str(job.project_id), "count": len(created)},
        organisation_id=job.organisation_id,
    )
    for sheet in created:
        if sheet.owner_id:
            emit(
                "sheet.assigned",
                {
                    "sheet_id": str(sheet.pk),
                    "user_id": str(sheet.owner_id),
                    "previous_user_id": None,
                },
                organisation_id=job.organisation_id,
            )
    return job


@transaction.atomic
def undo_import(actor: User, job: ImportJob) -> ImportJob:
    require(actor, Action.EDIT_SHEETS, project=job.project)
    job = ImportJob.objects.select_for_update().get(pk=job.pk)
    if job.state != ImportState.COMMITTED:
        raise ConflictError("not_undoable", "Only a committed import can be undone.")
    if job.undo_until is None or timezone.now() > job.undo_until:
        raise ConflictError("undo_expired", "Imports can be undone for 24 hours only.")
    sheets = list(Sheet.objects.filter(pk__in=job.created_sheet_ids))
    moved = [s.number for s in sheets if s.state != SheetState.PLANNED or s.revisions.exists()]
    if moved:
        raise ConflictError(
            "sheets_moved",
            f"{len(moved)} imported sheets have moved on ({', '.join(moved[:5])}"
            f"{'…' if len(moved) > 5 else ''}). The import can no longer be undone.",
            sheets=moved,
        )
    for sheet in sheets:
        record_change(
            sheet, actor=actor, verb="sheet.import_undone", before=snapshot(sheet), after=None
        )
    Sheet.objects.filter(pk__in=[s.pk for s in sheets]).delete()
    job.state = ImportState.UNDONE
    job.save(update_fields=["state", "updated_at"])
    record(
        organisation_id=job.organisation_id,
        actor=actor,
        verb="import.undone",
        entity_type="register.importjob",
        entity_id=job.pk,
        entity_label=job.original_name,
        project_id=job.project_id,
    )
    return job


def template_workbook() -> bytes:
    """Template for registers and onboarding (F-05, F-54). Onboarding sheets are added in S6."""
    from .onboarding_template import build_template

    return build_template()
